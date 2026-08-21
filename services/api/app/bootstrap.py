from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import tempfile
import time
import urllib.error
import urllib.request
import uuid
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable, Iterator, Sequence

import duckdb

from app.config import (
    LOCAL_CLEANED_XLSX_BYTES,
    LOCAL_CLEANED_XLSX_SHA256,
    Settings,
    get_settings,
)
from app.constants import (
    CATEGORY_TITLES,
    REQUIRED_COLUMNS,
    SESSION_CONSISTENCY_COLUMNS,
    SESSION_STATUSES,
    TRY_STATUSES,
    VERIFY_TYPES,
)


class BootstrapError(RuntimeError):
    """Base class for deterministic bootstrap failures."""


class SourceUnavailableError(BootstrapError):
    """The requested source could not be acquired."""


class DataContractError(BootstrapError):
    """The source violates a mandatory data-contract invariant."""

    def __init__(self, message: str, issues: Sequence[dict[str, object]] | None = None):
        super().__init__(message)
        self.issues = list(issues or [])


XLSX_MAX_DATA_ROWS = 1_048_575
EMBEDDED_DEMO_VERSION = "embedded-demo-calibrated-v2"
EMBEDDED_DEMO_SHA256 = hashlib.sha256(
    EMBEDDED_DEMO_VERSION.encode("utf-8")
).hexdigest()


def _has_spreadsheet_ceiling_risk(source_kind: str, attempt_rows: int) -> bool:
    return (
        source_kind in {"cleaned_xlsx", "cleaned_markdown"}
        and attempt_rows == XLSX_MAX_DATA_ROWS
    )


@dataclass(frozen=True, slots=True)
class SourceInfo:
    path: Path | None
    kind: str
    observed_sha256: str
    checksum_status: str
    source_reference: str
    byte_size: int | None
    partial_data: bool


@dataclass(frozen=True, slots=True)
class BootstrapResult:
    database_path: str
    build_id: str
    source_kind: str
    source_sha256: str
    checksum_status: str
    source_reference: str
    source_bytes: int | None
    attempt_rows: int
    session_rows: int
    customer_rows: int
    merchant_rows: int
    data_start: str
    data_end: str
    partial_data: bool
    generated_at: str


def _utc_now() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _sha256_file(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
            size += len(chunk)
    return digest.hexdigest(), size


def _sha256_directory(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    files = sorted(path.glob("records_*.md"))
    if not files:
        raise SourceUnavailableError(f"No cleaned Markdown record files found in {path}")
    for item in files:
        item_hash, item_size = _sha256_file(item)
        digest.update(item.name.encode("utf-8"))
        digest.update(item_hash.encode("ascii"))
        size += item_size
    return digest.hexdigest(), size


def _checksum_status(observed: str, expected: str | None) -> str:
    if not expected:
        return "observed-unpinned"
    if observed != expected.lower():
        raise SourceUnavailableError(
            f"Dataset SHA-256 mismatch: expected {expected.lower()}, observed {observed}"
        )
    return "verified"


def _download_dataset(settings: Settings, *, require_checksum: bool) -> SourceInfo:
    if require_checksum and not settings.dataset_sha256:
        raise SourceUnavailableError(
            "DATASET_SHA256 is required for non-demo official URL acquisition."
        )
    settings.data_dir.mkdir(parents=True, exist_ok=True)
    destination = settings.data_dir / "challenge_data.csv.gz"
    if destination.exists():
        observed, size = _sha256_file(destination)
        return SourceInfo(
            destination,
            "official_csv_gz_cache",
            observed,
            _checksum_status(observed, settings.dataset_sha256),
            settings.dataset_url,
            size,
            False,
        )

    partial = destination.with_suffix(destination.suffix + ".part")
    last_error: Exception | None = None
    for attempt in range(1, 5):
        try:
            request = urllib.request.Request(
                settings.dataset_url,
                headers={"User-Agent": "nabz-zarin-bootstrap/0.1"},
            )
            digest = hashlib.sha256()
            size = 0
            with urllib.request.urlopen(request, timeout=20) as response, partial.open(
                "wb"
            ) as output:
                while chunk := response.read(1024 * 1024):
                    output.write(chunk)
                    digest.update(chunk)
                    size += len(chunk)
            observed = digest.hexdigest()
            status = _checksum_status(observed, settings.dataset_sha256)
            os.replace(partial, destination)
            return SourceInfo(
                destination,
                "official_csv_gz_download",
                observed,
                status,
                settings.dataset_url,
                size,
                False,
            )
        except (OSError, urllib.error.URLError, TimeoutError) as exc:
            last_error = exc
            partial.unlink(missing_ok=True)
            if attempt < 4:
                time.sleep(min(2 ** (attempt - 1), 4))
    raise SourceUnavailableError(
        f"Official dataset could not be downloaded after 4 attempts: {last_error}"
    )


def _source_from_path(path: Path, expected_hash: str | None) -> SourceInfo:
    if not path.exists():
        raise SourceUnavailableError(f"Dataset path does not exist: {path}")
    if path.is_dir():
        observed, size = _sha256_directory(path)
        kind = "cleaned_markdown"
    else:
        observed, size = _sha256_file(path)
        suffixes = "".join(path.suffixes).lower()
        if suffixes.endswith(".csv.gz") or path.suffix.lower() == ".gz":
            kind = "csv_gz"
        elif path.suffix.lower() == ".csv":
            kind = "csv"
        elif path.suffix.lower() == ".xlsx":
            kind = "cleaned_xlsx"
        else:
            raise SourceUnavailableError(
                f"Unsupported dataset format for {path}; use CSV, CSV.GZ, XLSX, or a cleaned Markdown directory."
            )
    effective_expected = expected_hash
    if (
        effective_expected is None
        and path.is_file()
        and path.name == "challenge_data_cleaned.xlsx"
        and size == LOCAL_CLEANED_XLSX_BYTES
    ):
        effective_expected = LOCAL_CLEANED_XLSX_SHA256
    return SourceInfo(
        path,
        kind,
        observed,
        _checksum_status(observed, effective_expected),
        str(path),
        size,
        False,
    )


def resolve_source(
    settings: Settings,
    *,
    source: Path | None = None,
    download: bool = False,
    require_download_checksum: bool = False,
) -> SourceInfo:
    if settings.force_demo_data:
        return _embedded_demo_source()
    if source is not None:
        return _source_from_path(source.resolve(), settings.dataset_sha256)
    if settings.dataset_path is not None:
        return _source_from_path(settings.dataset_path, settings.dataset_sha256)

    # Local full-data fallbacks are preferred for reproducible development.
    workbook = settings.workspace_root / "challenge_data_cleaned.xlsx"
    if workbook.exists():
        return _source_from_path(workbook, None)
    markdown = settings.workspace_root / "outputs" / "challenge_data_cleaned"
    if markdown.exists():
        return _source_from_path(markdown, None)
    if download:
        return _download_dataset(settings, require_checksum=require_download_checksum)
    raise SourceUnavailableError("No local dataset was found and download was not requested.")


def _sql_string(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _create_empty_raw_table(connection: duckdb.DuckDBPyConnection) -> None:
    definitions = ", ".join(f'"{column}" VARCHAR' for column in REQUIRED_COLUMNS)
    connection.execute(f"CREATE TABLE source_raw ({definitions})")


def _load_csv(connection: duckdb.DuckDBPyConnection, path: Path) -> None:
    connection.execute(
        """
        CREATE TABLE source_raw AS
        SELECT * FROM read_csv(
            ?, header = true, all_varchar = true, auto_detect = true,
            nullstr = ['null', 'NULL'], sample_size = -1, strict_mode = true
        )
        """,
        [str(path)],
    )


def _bulk_load_rows(
    connection: duckdb.DuckDBPyConnection,
    rows: Iterable[Sequence[object | None]],
    *,
    temporary_directory: Path,
) -> None:
    """Stage non-CSV sources once, then let DuckDB perform a vectorized scan.

    This deliberately avoids ``executemany`` for the million-row workbook: on
    Windows that path can take hours and historically left an open temporary
    SQLite/DB handle behind. The named file is closed before DuckDB opens it and
    is removed in ``finally`` on every platform.
    """

    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=".zarinpal-normalized-", suffix=".csv", dir=temporary_directory
    )
    os.close(file_descriptor)
    staged = Path(temporary_name)
    try:
        with staged.open("w", encoding="utf-8", newline="") as output:
            writer = csv.writer(output, lineterminator="\n")
            writer.writerow(REQUIRED_COLUMNS)
            writer.writerows(rows)
        _load_csv(connection, staged)
    finally:
        staged.unlink(missing_ok=True)


def _load_xlsx(connection: duckdb.DuckDBPyConnection, path: Path) -> None:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise SourceUnavailableError("XLSX fallback requires openpyxl") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value).strip() if value is not None else "" for value in next(rows))
        missing = sorted(set(REQUIRED_COLUMNS) - set(headers))
        if missing:
            raise DataContractError(f"Required columns are absent: {', '.join(missing)}")
        indices = [headers.index(column) for column in REQUIRED_COLUMNS]
        def normalized_rows() -> Iterator[tuple[str | None, ...]]:
            for row in rows:
                normalized: list[str | None] = []
                for index in indices:
                    value = row[index] if index < len(row) else None
                    if value is None:
                        normalized.append(None)
                    elif isinstance(value, datetime):
                        normalized.append(value.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3])
                    else:
                        normalized.append(str(value))
                yield tuple(normalized)

        _bulk_load_rows(
            connection,
            normalized_rows(),
            temporary_directory=Path(tempfile.gettempdir()),
        )
    finally:
        workbook.close()


def _markdown_rows(path: Path) -> Iterator[tuple[str | None, ...]]:
    for item in sorted(path.glob("records_*.md")):
        with item.open("r", encoding="utf-8") as stream:
            data_started = False
            for raw_line in stream:
                line = raw_line.rstrip("\r\n")
                if not data_started:
                    if line.startswith("|---"):
                        data_started = True
                    continue
                if not line.startswith("|"):
                    continue
                values = [value.strip() for value in line.strip().strip("|").split("|")]
                if len(values) != len(REQUIRED_COLUMNS):
                    raise DataContractError(
                        f"Malformed Markdown record in {item.name}: expected {len(REQUIRED_COLUMNS)} fields, got {len(values)}"
                    )
                yield tuple(None if value == "null" else value.replace("\\|", "|") for value in values)


def _load_markdown(connection: duckdb.DuckDBPyConnection, path: Path) -> None:
    _bulk_load_rows(
        connection,
        _markdown_rows(path),
        temporary_directory=Path(tempfile.gettempdir()),
    )


def _load_demo(connection: duckdb.DuckDBPyConnection) -> None:
    """Create a calibrated, pseudonymous fallback without claiming full fidelity.

    The three featured merchants preserve the challenge-wide session, verified,
    retry, rescue and (where known) revenue totals used by the demo narrative.
    All row-level paths are synthetic and the manifest marks the database partial.
    """

    connection.execute(
        """
        CREATE TABLE source_raw AS
        WITH base_configs AS (
            SELECT * FROM (VALUES
                ('M43', 29483, 20340, 190, 7922, 6057, 500, 56610001, 'کیف و کفش فروشی', 'T196', 23448530700::BIGINT, 9745, 3692),
                ('M31', 194600, 37437, 0, 3346, 485, 1000, 82410000, 'مراکز آموزشی مجازی', 'T99', 64151420000::BIGINT, 30740, 5268),
                ('M156', 21776, 12268, 288, 1372, 575, 200, 59770001, 'فروشگاه لوازم آرایشی و بهداشتی', 'T207', 731158025930::BIGINT, 10469, 1226)
            ) AS t(merchant_key, total_sessions, verified_sessions, paid_sessions, retry_sessions, rescued_sessions, no_attempt_sessions, category_id, category_title, terminal_key, verified_revenue, buyer_cards, repeat_cards)
        ),
        sessions AS (
            SELECT c.*, i,
                CASE
                    WHEN i <= rescued_sessions THEN 'Verified'
                    WHEN i > retry_sessions AND i <= retry_sessions + verified_sessions - rescued_sessions THEN 'Verified'
                    WHEN i > retry_sessions + verified_sessions - rescued_sessions
                         AND i <= retry_sessions + verified_sessions - rescued_sessions + paid_sessions THEN 'Paid'
                    ELSE 'Failed'
                END AS final_status,
                CASE
                    WHEN i <= rescued_sessions THEN i
                    WHEN i > retry_sessions AND i <= retry_sessions + verified_sessions - rescued_sessions THEN rescued_sessions + i - retry_sessions
                    ELSE NULL
                END AS verified_ordinal
            FROM base_configs c, LATERAL range(1, c.total_sessions + 1) AS ids(i)
        ),
        session_values AS (
            SELECT *,
                CASE
                    WHEN verified_revenue IS NOT NULL AND final_status = 'Verified'
                    THEN floor(verified_revenue / verified_sessions)::BIGINT
                         + CASE WHEN verified_ordinal <= verified_revenue % verified_sessions THEN 1 ELSE 0 END
                    ELSE 1000000::BIGINT + (i % 7) * 100000
                END AS amount_value,
                TIMESTAMP '2026-01-01 00:00:00'
                    + ((i - 1) % CASE WHEN merchant_key = 'M43' THEN 59 ELSE 120 END)
                      * INTERVAL '1 day'
                    + (i % 86400) * INTERVAL '1 second' AS created_value
            FROM sessions
        ),
        attempts AS (
            SELECT *,
                CASE WHEN i > total_sessions - no_attempt_sessions THEN 0 ELSE 1 END AS seq,
                CASE
                    WHEN i > total_sessions - no_attempt_sessions THEN 'NoAttempt'
                    WHEN i <= retry_sessions THEN 'InBank'
                    ELSE final_status
                END AS attempt_status
            FROM session_values
            UNION ALL
            SELECT *, 2 AS seq, final_status AS attempt_status
            FROM session_values WHERE i <= retry_sessions
        )
        SELECT
            merchant_key || '-' || i::VARCHAR AS session_key,
            seq::VARCHAR AS try_seq,
            terminal_key,
            merchant_key,
            category_id::VARCHAR AS category_id,
            category_title,
            amount_value::VARCHAR AS amount,
            greatest(1000, floor(amount_value * 0.008))::BIGINT::VARCHAR AS adjusted_fee,
            final_status AS session_status,
            attempt_status AS try_status,
            CASE WHEN seq > 0 AND attempt_status = 'Failed' THEN 'PSP-03:demo' ELSE NULL END AS switch_response_code,
            CASE WHEN seq > 0 THEN CASE WHEN seq = 1 THEN 'PSP-03' ELSE 'PSP-04' END ELSE NULL END AS psp_code,
            CASE WHEN final_status IN ('Verified', 'Paid') AND seq = CASE WHEN i <= retry_sessions THEN 2 ELSE 1 END THEN 'BANK-DEMO' ELSE NULL END AS issuer_bank_code,
            CASE
                WHEN final_status NOT IN ('Verified', 'Paid') OR seq <> CASE WHEN i <= retry_sessions THEN 2 ELSE 1 END THEN NULL
                WHEN verified_ordinal IS NOT NULL AND verified_ordinal <= repeat_cards * 2
                    THEN 'CARD-' || merchant_key || '-' || ceil(verified_ordinal / 2.0)::BIGINT::VARCHAR
                WHEN verified_ordinal IS NOT NULL AND verified_ordinal <= repeat_cards * 2 + (buyer_cards - repeat_cards)
                    THEN 'CARD-' || merchant_key || '-' || (repeat_cards + verified_ordinal - repeat_cards * 2)::BIGINT::VARCHAR
                WHEN verified_ordinal IS NOT NULL
                    THEN 'CARD-' || merchant_key || '-' || (1 + ((verified_ordinal - 1) % repeat_cards))::BIGINT::VARCHAR
                ELSE 'CARD-' || merchant_key || '-' || (1 + ((i - 1) % greatest(buyer_cards, 1)))::VARCHAR
            END AS payer_card_key,
            'Automated' AS verify_type,
            CASE WHEN seq > 0 THEN (70 + (i % 230))::VARCHAR ELSE NULL END AS init_time_ms,
            CASE WHEN final_status = 'Verified' AND seq = CASE WHEN i <= retry_sessions THEN 2 ELSE 1 END THEN (60 + (i % 190))::VARCHAR ELSE NULL END AS verify_time_ms,
            strftime(created_value, '%Y-%m-%d %H:%M:%S.%g') AS created_at,
            CASE WHEN seq > 0 THEN strftime(created_value + seq * INTERVAL '1 second', '%Y-%m-%d %H:%M:%S.%g') ELSE NULL END AS try_created_at,
            CASE WHEN final_status = 'Verified' THEN strftime(created_value + INTERVAL '90 seconds', '%Y-%m-%d %H:%M:%S.%g') ELSE NULL END AS verified_at,
            CASE WHEN final_status = 'Verified' THEN strftime(created_value + INTERVAL '88 seconds', '%Y-%m-%d %H:%M:%S.%g') ELSE NULL END AS settled_at,
            strftime(created_value + INTERVAL '30 minutes', '%Y-%m-%d %H:%M:%S.%g') AS expire_in
        FROM attempts
        """
    )


def _load_source(connection: duckdb.DuckDBPyConnection, source: SourceInfo) -> None:
    if source.kind == "embedded_demo":
        _load_demo(connection)
    elif source.kind in {"csv", "csv_gz", "official_csv_gz_cache", "official_csv_gz_download"}:
        assert source.path is not None
        _load_csv(connection, source.path)
    elif source.kind == "cleaned_xlsx":
        assert source.path is not None
        _load_xlsx(connection, source.path)
    elif source.kind == "cleaned_markdown":
        assert source.path is not None
        _load_markdown(connection, source.path)
    else:  # pragma: no cover - resolve_source prevents this
        raise SourceUnavailableError(f"Unsupported source kind: {source.kind}")


def _validate_columns(connection: duckdb.DuckDBPyConnection) -> None:
    columns = {row[1] for row in connection.execute("PRAGMA table_info('source_raw')").fetchall()}
    missing = sorted(set(REQUIRED_COLUMNS) - columns)
    if missing:
        raise DataContractError(f"Required columns are absent: {', '.join(missing)}")


def _count(connection: duckdb.DuckDBPyConnection, sql: str) -> int:
    return int(connection.execute(sql).fetchone()[0])


def _validate_and_normalize(connection: duckdb.DuckDBPyConnection) -> None:
    _validate_columns(connection)
    issues: list[dict[str, object]] = []

    checks = {
        "missing_required_session_fields": """
            SELECT count(*) FROM source_raw WHERE
                nullif(trim(session_key), '') IS NULL OR
                nullif(trim(terminal_key), '') IS NULL OR
                nullif(trim(merchant_key), '') IS NULL OR
                nullif(trim(category_id), '') IS NULL OR
                nullif(trim(category_title), '') IS NULL OR
                nullif(trim(amount), '') IS NULL OR
                nullif(trim(adjusted_fee), '') IS NULL OR
                nullif(trim(session_status), '') IS NULL OR
                nullif(trim(verify_type), '') IS NULL OR
                nullif(trim(created_at), '') IS NULL OR
                nullif(trim(expire_in), '') IS NULL
        """,
        "invalid_numeric_fields": """
            SELECT count(*) FROM source_raw WHERE
                try_cast(trim(try_seq) AS INTEGER) IS NULL OR
                try_cast(trim(category_id) AS BIGINT) IS NULL OR
                try_cast(trim(amount) AS BIGINT) IS NULL OR
                try_cast(trim(adjusted_fee) AS BIGINT) IS NULL OR
                try_cast(trim(amount) AS BIGINT) < 0 OR
                try_cast(trim(adjusted_fee) AS BIGINT) < 0
        """,
        "invalid_required_timestamps": """
            SELECT count(*) FROM source_raw WHERE
                try_cast(trim(created_at) AS TIMESTAMP) IS NULL OR
                try_cast(trim(expire_in) AS TIMESTAMP) IS NULL
        """,
    }
    for name, sql in checks.items():
        issue_count = _count(connection, sql)
        issues.append({"check_name": name, "severity": "error", "issue_count": issue_count})

    if any(int(issue["issue_count"]) > 0 for issue in issues):
        raise DataContractError("Source failed type and required-value validation.", issues)

    category_case = "CASE try_cast(trim(category_id) AS BIGINT) " + " ".join(
        f"WHEN {category_id} THEN {_sql_string(title)}"
        for category_id, title in CATEGORY_TITLES.items()
    ) + " ELSE trim(category_title) END"

    connection.execute(
        f"""
        CREATE TABLE attempt_fact AS
        SELECT
            trim(session_key)::VARCHAR AS session_key,
            try_cast(trim(try_seq) AS INTEGER) AS try_seq,
            trim(terminal_key)::VARCHAR AS terminal_key,
            trim(merchant_key)::VARCHAR AS merchant_key,
            try_cast(trim(category_id) AS BIGINT) AS category_id,
            {category_case}::VARCHAR AS category_title,
            try_cast(trim(amount) AS BIGINT) AS amount,
            try_cast(trim(adjusted_fee) AS BIGINT) AS adjusted_fee,
            trim(session_status)::VARCHAR AS session_status,
            trim(try_status)::VARCHAR AS try_status,
            nullif(trim(switch_response_code), '')::VARCHAR AS switch_response_code,
            nullif(trim(psp_code), '')::VARCHAR AS psp_code,
            nullif(trim(issuer_bank_code), '')::VARCHAR AS issuer_bank_code,
            nullif(trim(payer_card_key), '')::VARCHAR AS payer_card_key,
            trim(verify_type)::VARCHAR AS verify_type,
            try_cast(nullif(trim(init_time_ms), '') AS INTEGER) AS init_time_ms,
            try_cast(nullif(trim(verify_time_ms), '') AS INTEGER) AS verify_time_ms,
            try_cast(trim(created_at) AS TIMESTAMP) AS created_at,
            try_cast(nullif(trim(try_created_at), '') AS TIMESTAMP) AS try_created_at,
            try_cast(nullif(trim(verified_at), '') AS TIMESTAMP) AS verified_at,
            try_cast(nullif(trim(settled_at), '') AS TIMESTAMP) AS settled_at,
            try_cast(trim(expire_in) AS TIMESTAMP) AS expire_in
        FROM source_raw
        """
    )
    connection.execute("DROP TABLE source_raw")

    invariant_checks = {
        "duplicate_attempt_key": """
            SELECT count(*) FROM (
                SELECT session_key, try_seq FROM attempt_fact
                GROUP BY session_key, try_seq HAVING count(*) > 1
            )
        """,
        "invalid_try_seq": "SELECT count(*) FROM attempt_fact WHERE try_seq < 0",
        "invalid_session_status": (
            "SELECT count(*) FROM attempt_fact WHERE session_status NOT IN ("
            + ",".join(_sql_string(value) for value in SESSION_STATUSES)
            + ")"
        ),
        "invalid_try_status": (
            "SELECT count(*) FROM attempt_fact WHERE try_status NOT IN ("
            + ",".join(_sql_string(value) for value in TRY_STATUSES)
            + ")"
        ),
        "invalid_verify_type": (
            "SELECT count(*) FROM attempt_fact WHERE verify_type NOT IN ("
            + ",".join(_sql_string(value) for value in VERIFY_TYPES)
            + ")"
        ),
        "invalid_no_attempt_state": """
            SELECT count(*) FROM attempt_fact
            WHERE (try_seq = 0 AND try_status <> 'NoAttempt')
               OR (try_seq > 0 AND try_status = 'NoAttempt')
               OR (try_seq = 0 AND (psp_code IS NOT NULL OR try_created_at IS NOT NULL))
        """,
    }
    for name, sql in invariant_checks.items():
        issue_count = _count(connection, sql)
        issues.append({"check_name": name, "severity": "error", "issue_count": issue_count})

    union_queries = []
    for column in SESSION_CONSISTENCY_COLUMNS:
        union_queries.append(
            f"""
            SELECT session_key, {_sql_string(column)} AS column_name
            FROM attempt_fact
            GROUP BY session_key
            HAVING count(DISTINCT coalesce(cast(\"{column}\" AS VARCHAR), '__NULL__')) > 1
            """
        )
    connection.execute(
        """
        CREATE TABLE session_consistency_report AS
        SELECT session_key, string_agg(column_name, ',' ORDER BY column_name) AS conflicting_columns
        FROM (
        """
        + " UNION ALL ".join(union_queries)
        + ") conflicts GROUP BY session_key"
    )
    conflict_count = _count(connection, "SELECT count(*) FROM session_consistency_report")
    issues.append(
        {
            "check_name": "inconsistent_session_fields",
            "severity": "error",
            "issue_count": conflict_count,
        }
    )

    connection.execute(
        """
        CREATE TABLE data_quality_report (
            check_name VARCHAR, severity VARCHAR, issue_count BIGINT, details VARCHAR
        )
        """
    )
    connection.executemany(
        "INSERT INTO data_quality_report VALUES (?, ?, ?, ?)",
        [
            (
                str(issue["check_name"]),
                str(issue["severity"]),
                int(issue["issue_count"]),
                "Mandatory data-contract validation",
            )
            for issue in issues
        ],
    )
    failures = [issue for issue in issues if int(issue["issue_count"]) > 0]
    if failures:
        raise DataContractError("Source violates mandatory data-contract invariants.", failures)


def _materialize(connection: duckdb.DuckDBPyConnection) -> None:
    connection.execute(
        """
        CREATE TABLE session_fact AS
        SELECT
            session_key,
            any_value(terminal_key) AS terminal_key,
            any_value(merchant_key) AS merchant_key,
            any_value(category_id) AS category_id,
            any_value(category_title) AS category_title,
            any_value(amount) AS amount,
            any_value(adjusted_fee) AS adjusted_fee,
            any_value(session_status) AS session_status,
            any_value(verify_type) AS verify_type,
            any_value(created_at) AS created_at,
            any_value(verified_at) AS verified_at,
            any_value(settled_at) AS settled_at,
            any_value(expire_in) AS expire_in,
            count(*) FILTER (WHERE try_seq > 0)::INTEGER AS attempt_count,
            bool_or(try_seq > 0) AS has_attempt,
            coalesce(max(try_seq) FILTER (WHERE try_seq > 0), 0)::INTEGER AS max_try_seq,
            count(DISTINCT psp_code) FILTER (WHERE try_seq > 0 AND psp_code IS NOT NULL)::INTEGER AS distinct_psp_count,
            string_agg(try_status, ' → ' ORDER BY try_seq) FILTER (WHERE try_seq > 0) AS attempt_status_path,
            string_agg(psp_code, ' → ' ORDER BY try_seq) FILTER (WHERE try_seq > 0 AND psp_code IS NOT NULL) AS psp_path,
            arg_max(try_status, try_seq) FILTER (WHERE try_seq > 0) AS final_try_status,
            arg_max(payer_card_key, try_seq) FILTER (WHERE payer_card_key IS NOT NULL) AS payer_card_key,
            arg_max(issuer_bank_code, try_seq) FILTER (WHERE issuer_bank_code IS NOT NULL) AS issuer_bank_code,
            bool_or(try_seq > 0 AND try_status IN ('InBank', 'Paid', 'Verified')) AS reached_inbank,
            bool_or(try_status IN ('Paid', 'Verified')) OR any_value(session_status) IN ('Paid', 'Verified') AS reached_paid,
            any_value(session_status) = 'Verified' AS is_verified,
            any_value(session_status) = 'Paid' AS is_paid_not_verified,
            CASE WHEN any_value(verified_at) IS NOT NULL
                 THEN date_diff('millisecond', any_value(created_at), any_value(verified_at))
                 ELSE NULL END AS time_to_verify_ms,
            CASE WHEN any_value(settled_at) IS NOT NULL
                 THEN date_diff('millisecond', any_value(created_at), any_value(settled_at))
                 ELSE NULL END AS time_to_settle_ms
        FROM attempt_fact
        GROUP BY session_key
        """
    )
    connection.execute("CREATE UNIQUE INDEX session_fact_pk ON session_fact(session_key)")
    connection.execute(
        "CREATE UNIQUE INDEX attempt_fact_pk ON attempt_fact(session_key, try_seq)"
    )
    connection.execute(
        "CREATE INDEX session_merchant_date ON session_fact(merchant_key, created_at)"
    )
    connection.execute(
        "CREATE INDEX attempt_merchant_date ON attempt_fact(merchant_key, created_at)"
    )

    connection.execute(
        """
        CREATE TABLE merchant_customer_fact AS
        SELECT
            merchant_key,
            payer_card_key,
            count(*)::INTEGER AS verified_session_count,
            sum(amount)::BIGINT AS verified_revenue,
            min(created_at) AS first_verified_at,
            max(created_at) AS last_verified_at,
            count(*) >= 2 AS is_observed_repeat
        FROM session_fact
        WHERE is_verified AND payer_card_key IS NOT NULL
        GROUP BY merchant_key, payer_card_key
        """
    )
    connection.execute(
        "CREATE UNIQUE INDEX merchant_customer_pk ON merchant_customer_fact(merchant_key, payer_card_key)"
    )

    connection.execute(
        """
        CREATE TABLE merchant_profile AS
        WITH merchant_totals AS (
            SELECT merchant_key, any_value(category_id) AS category_id,
                   any_value(category_title) AS category_title,
                   count(*)::BIGINT AS session_count
            FROM session_fact GROUP BY merchant_key
        )
        SELECT *, ntile(3) OVER (PARTITION BY category_id ORDER BY session_count, merchant_key) AS volume_tercile
        FROM merchant_totals
        """
    )
    connection.execute(
        """
        CREATE TABLE merchant_daily AS
        SELECT
            merchant_key,
            cast(created_at AS DATE) AS activity_date,
            count(*)::BIGINT AS session_count,
            count(*) FILTER (WHERE is_verified)::BIGINT AS verified_session_count,
            coalesce(sum(amount) FILTER (WHERE is_verified), 0)::BIGINT AS verified_revenue,
            count(*) FILTER (WHERE is_paid_not_verified)::BIGINT AS paid_not_verified_count,
            count(*) FILTER (WHERE attempt_count = 0)::BIGINT AS no_attempt_count,
            count(*) FILTER (WHERE attempt_count > 1)::BIGINT AS retry_session_count,
            count(*) FILTER (WHERE attempt_count > 1 AND is_verified)::BIGINT AS rescued_session_count
        FROM session_fact
        GROUP BY merchant_key, cast(created_at AS DATE)
        """
    )

    missing_columns = (
        "issuer_bank_code",
        "payer_card_key",
        "psp_code",
        "try_created_at",
        "verified_at",
        "settled_at",
        "init_time_ms",
        "verify_time_ms",
    )
    missing_queries = []
    for column in missing_columns:
        missing_queries.append(
            f"""
            SELECT session_status AS lifecycle_status, {_sql_string(column)} AS column_name,
                   count(*) FILTER (WHERE \"{column}\" IS NULL)::BIGINT AS null_count,
                   count(*)::BIGINT AS row_count,
                   count(*) FILTER (WHERE \"{column}\" IS NULL)::DOUBLE / nullif(count(*), 0) AS null_rate
            FROM attempt_fact GROUP BY session_status
            """
        )
    connection.execute(
        "CREATE TABLE missingness_report AS " + " UNION ALL ".join(missing_queries)
    )
    connection.execute(
        """
        CREATE TABLE concentration_report AS
        WITH merchant_counts AS (
            SELECT merchant_key, count(*)::BIGINT AS session_count FROM session_fact GROUP BY merchant_key
        ), ranked AS (
            SELECT *, row_number() OVER (ORDER BY session_count DESC, merchant_key) AS rank,
                   sum(session_count) OVER () AS all_sessions
            FROM merchant_counts
        )
        SELECT merchant_key, session_count, rank,
               session_count::DOUBLE / nullif(all_sessions, 0) AS session_share
        FROM ranked ORDER BY rank
        """
    )


def _write_manifest(
    connection: duckdb.DuckDBPyConnection,
    source: SourceInfo,
    build_id: str,
    generated_at: str,
    database_path: Path,
) -> BootstrapResult:
    counts = connection.execute(
        """
        SELECT
            (SELECT count(*) FROM attempt_fact),
            (SELECT count(*) FROM session_fact),
            (SELECT count(*) FROM merchant_customer_fact),
            (SELECT count(*) FROM merchant_profile),
            (SELECT min(cast(created_at AS DATE)) FROM session_fact)::VARCHAR,
            (SELECT max(cast(created_at AS DATE)) FROM session_fact)::VARCHAR
        """
    ).fetchone()
    attempt_rows = int(counts[0])
    spreadsheet_ceiling_risk = _has_spreadsheet_ceiling_risk(
        source.kind, attempt_rows
    )
    result = BootstrapResult(
        database_path=str(database_path),
        build_id=build_id,
        source_kind=source.kind,
        source_sha256=source.observed_sha256,
        checksum_status=source.checksum_status,
        source_reference=source.source_reference,
        source_bytes=source.byte_size,
        attempt_rows=attempt_rows,
        session_rows=int(counts[1]),
        customer_rows=int(counts[2]),
        merchant_rows=int(counts[3]),
        data_start=str(counts[4]),
        data_end=str(counts[5]),
        # Excel permits 1,048,576 rows including the header. A repaired XLSX
        # (or Markdown exported from it) landing exactly on that ceiling cannot
        # establish challenge-wide completeness, even if every retained row is
        # internally valid. Only the pinned official CSV/GZ can remove this flag.
        partial_data=source.partial_data or spreadsheet_ceiling_risk,
        generated_at=generated_at,
    )
    connection.execute("CREATE TABLE build_manifest (key VARCHAR PRIMARY KEY, value VARCHAR)")
    values = asdict(result)
    connection.executemany(
        "INSERT INTO build_manifest VALUES (?, ?)",
        [(key, json.dumps(value, ensure_ascii=False) if not isinstance(value, str) else value) for key, value in values.items()],
    )
    return result


def _embedded_demo_source() -> SourceInfo:
    return SourceInfo(
        path=None,
        kind="embedded_demo",
        observed_sha256=EMBEDDED_DEMO_SHA256,
        checksum_status="embedded-versioned",
        source_reference=EMBEDDED_DEMO_VERSION,
        byte_size=None,
        partial_data=True,
    )


def bootstrap_database(
    settings: Settings | None = None,
    *,
    database_path: Path | None = None,
    source: Path | None = None,
    download: bool = False,
    allow_demo_fallback: bool | None = None,
) -> BootstrapResult:
    settings = settings or get_settings()
    target = (database_path or settings.database_path).resolve()
    allow_demo = settings.allow_demo_fallback if allow_demo_fallback is None else allow_demo_fallback
    target.parent.mkdir(parents=True, exist_ok=True)

    try:
        source_info = resolve_source(
            settings,
            source=source,
            download=download,
            require_download_checksum=not allow_demo,
        )
    except SourceUnavailableError:
        if not allow_demo:
            raise
        source_info = _embedded_demo_source()

    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{target.stem}.building-", suffix=".duckdb", dir=target.parent
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    # DuckDB expects either a valid database or a path that does not exist.
    temporary_path.unlink(missing_ok=True)
    connection: duckdb.DuckDBPyConnection | None = None
    generated_at = _utc_now()
    build_id = str(uuid.uuid4())
    try:
        connection = duckdb.connect(str(temporary_path))
        connection.execute("SET TimeZone='Asia/Tehran'")
        connection.execute("SET preserve_insertion_order=false")
        _load_source(connection, source_info)
        _validate_and_normalize(connection)
        _materialize(connection)
        result = _write_manifest(connection, source_info, build_id, generated_at, target)
        connection.execute("CHECKPOINT")
        connection.close()
        connection = None
        os.replace(temporary_path, target)
        return BootstrapResult(**{**asdict(result), "database_path": str(target)})
    finally:
        if connection is not None:
            connection.close()
        temporary_path.unlink(missing_ok=True)
        Path(str(temporary_path) + ".wal").unlink(missing_ok=True)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build the ZarinPal analytical DuckDB")
    parser.add_argument("--source", type=Path, help="CSV/GZ/XLSX or cleaned Markdown directory")
    parser.add_argument("--database", type=Path, help="Output DuckDB path")
    parser.add_argument("--download", action="store_true", help="Download the official source if no local source exists")
    parser.add_argument("--allow-demo-fallback", action="store_true", help="Use the calibrated partial demo if acquisition fails")
    parser.add_argument("--json-output", type=Path, help="Optional manifest output path")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    result = bootstrap_database(
        database_path=args.database,
        source=args.source,
        download=args.download,
        allow_demo_fallback=args.allow_demo_fallback,
    )
    payload = json.dumps(asdict(result), ensure_ascii=False, indent=2)
    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(payload + "\n", encoding="utf-8")
    print(payload)
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
